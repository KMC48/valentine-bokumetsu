# -*- coding: utf-8 -*-
"""ストーリー台本（オープニング・中間×2 ×3周回、エンディング6種）を Word にまとめる。"""
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

data = json.load(open("scripts/_script_extracted.json", encoding="utf-8"))
stories = data["stories"]
endings = data["endings"]

LOOP_NAME = {1: "1周目：県立桜庭高等学校（探す）", 2: "2周目：私立聖蘭学院（見破る）", 3: "3周目：白鷺台学園（推理する）"}
SLOT_NAME = {"opening": "オープニング（朝ステージ前）", "interlude1": "中間ストーリー①（昼ステージ前）", "interlude2": "中間ストーリー②（放課後ステージ前）"}
RANK_NAME = {"S": "Sランク", "A": "Aランク", "B": "Bランク", "C": "Cランク", "D": "D以下", "special": "特殊エンディング"}

doc = Document()

# 既定フォントを日本語対応に。
style = doc.styles["Normal"]
style.font.name = "游ゴシック"
style.font.size = Pt(10.5)
rpr = style.element.get_or_add_rPr()
rFonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
if rFonts is None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rFonts = OxmlElement("w:rFonts")
    rpr.append(rFonts)
from docx.oxml.ns import qn

rFonts.set(qn("w:eastAsia"), "游ゴシック")

for sec in doc.sections:
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

title = doc.add_heading("バレンタイン撲滅委員会 ストーリー台本", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph("オープニング・中間ストーリー（周回1〜3）／エンディング（ランク別5種＋特殊）")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(11)
sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_page_break()

# ---------- 目次代わりの一覧 ----------
doc.add_heading("収録内容", level=1)
toc = doc.add_paragraph()
toc.add_run("■ ストーリーパート：全9シーン（3周回 × オープニング/中間①/中間②）\n").bold = True
for loop in (1, 2, 3):
    toc.add_run(f"　・{LOOP_NAME[loop]}\n")
toc.add_run("\n■ エンディング：全6種（S/A/B/C/D ランク＋特殊エンディング）\n").bold = True

doc.add_page_break()


def add_speaker_line(doc, speaker, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if speaker:
        run = p.add_run(f"【{speaker}】")
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        p.add_run("\n")
    run2 = p.add_run(text)
    run2.font.size = Pt(11)


# ---------- ストーリーパート ----------
doc.add_heading("ストーリーパート", level=1)

for loop in (1, 2, 3):
    doc.add_heading(LOOP_NAME[loop], level=2)
    scenes = [s for s in stories if s["loop"] == loop]
    order = ["opening", "interlude1", "interlude2"]
    scenes.sort(key=lambda s: order.index(s["slot"]))
    for scene in scenes:
        doc.add_heading(f"{SLOT_NAME.get(scene['slot'], scene['slot'])} 「{scene['title']}」", level=3)
        for line in scene["lines"]:
            add_speaker_line(doc, line["speaker"], line["text"])
        doc.add_paragraph()
    doc.add_page_break()

# ---------- エンディング ----------
doc.add_heading("エンディング", level=1)

rank_order = ["S", "A", "B", "C", "D", "special"]
endings_sorted = sorted(endings, key=lambda e: rank_order.index(e["rank"]) if e["rank"] in rank_order else 99)

for ending in endings_sorted:
    doc.add_heading(f"{RANK_NAME.get(ending['rank'], ending['rank'])} 「{ending['title']}」", level=2)
    for line in ending["lines"]:
        add_speaker_line(doc, line["speaker"], line["text"])
    doc.add_paragraph()

OUT_PATH = "ストーリー台本.docx"
doc.save(OUT_PATH)
print("saved:", OUT_PATH)
