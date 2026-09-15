# -*- coding: utf-8 -*-
"""観察メモ一覧（45人分）を Excel にまとめる。"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

rows = json.load(open("scripts/_students_extracted.json", encoding="utf-8"))

LOOP_NAME = {1: "1周目：県立桜庭高等学校", 2: "2周目：私立聖蘭学院", 3: "3周目：白鷺台学園"}
STAGE_NAME = {"morning": "朝：登校", "lunch": "昼休み", "afterSchool": "放課後"}
CHOCO_NAME = {
    "none": "無実",
    "friend": "友チョコ",
    "honmei": "本命チョコ",
    "honmei_special": "本命（告白直前）",
}
ITEM_NAME = {
    "none": "持ち物なし",
    "paper_bag": "紙袋",
    "decoy_bag": "紙袋（囮）",
    "lunch_box": "弁当箱",
    "book_case": "教科書ケース",
    "pouch": "ポーチ",
    "sports_bag": "部活バッグ",
}

HEADERS = [
    "周回", "ステージ", "ID", "名前", "チョコ種別", "怪しさ", "所持品", "囮フラグ",
    "観察メモ①", "観察メモ②", "観察メモ③",
    "セリフ", "仕草", "通報結果テキスト", "配置X(%)", "配置Y(%)",
]

wb = Workbook()
ws = wb.active
ws.title = "観察メモ一覧"

header_fill = PatternFill(start_color="2F4F6F", end_color="2F4F6F", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
no_fill = PatternFill(fill_type=None)
thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center")

for col, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.border = border
    c.alignment = center
ws.freeze_panes = "A2"

# 周回→ステージ→配置X の順で並べる（ゲーム内の出現順に近づける）。
rows_sorted = sorted(rows, key=lambda r: (r["loop"], ["morning", "lunch", "afterSchool"].index(r["stage"]), r["x"]))

for i, r in enumerate(rows_sorted):
    row_idx = i + 2
    hints = r["hints"] + [""] * 3
    values = [
        LOOP_NAME.get(r["loop"], r["loop"]),
        STAGE_NAME.get(r["stage"], r["stage"]),
        r["id"],
        r["name"],
        CHOCO_NAME.get(r["chocolateType"], r["chocolateType"]),
        r["suspicion"],
        ITEM_NAME.get(r["item"], r["item"]),
        "○" if r["isDecoy"] else "",
        hints[0], hints[1], hints[2],
        " / ".join(r["dialogue"]),
        " / ".join(r["behavior"]),
        r["result"] or "",
        r["x"],
        r["y"],
    ]
    fill = blue_fill if (i % 2 == 0) else no_fill
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col, value=v)
        cell.fill = fill
        cell.border = border
        cell.alignment = wrap_top

# 列幅の目安。
WIDTHS = [20, 10, 10, 14, 16, 8, 14, 6, 26, 26, 26, 30, 24, 32, 10, 10]
for col, w in enumerate(WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w

ws.row_dimensions[1].height = 22

# 周回ごとの集計シートも追加。
summary = wb.create_sheet("集計")
summary_headers = ["周回", "友チョコ", "本命チョコ", "本命（告白直前）", "無実", "合計"]
for col, h in enumerate(summary_headers, start=1):
    c = summary.cell(row=1, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.border = border
    c.alignment = center

for i, loop in enumerate([1, 2, 3]):
    loop_rows = [r for r in rows if r["loop"] == loop]
    counts = {k: sum(1 for r in loop_rows if r["chocolateType"] == k) for k in CHOCO_NAME}
    values = [
        LOOP_NAME[loop],
        counts["friend"],
        counts["honmei"],
        counts["honmei_special"],
        counts["none"],
        len(loop_rows),
    ]
    fill = blue_fill if (i % 2 == 0) else no_fill
    for col, v in enumerate(values, start=1):
        cell = summary.cell(row=i + 2, column=col, value=v)
        cell.fill = fill
        cell.border = border
        cell.alignment = center

for col, w in enumerate([24, 10, 12, 16, 8, 8], start=1):
    summary.column_dimensions[get_column_letter(col)].width = w

OUT_PATH = "観察メモ一覧.xlsx"
wb.save(OUT_PATH)
print("saved:", OUT_PATH, "rows:", len(rows_sorted))
